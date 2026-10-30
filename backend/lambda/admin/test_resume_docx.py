"""Word ingest and rewrite keep locked paragraphs and reject unsafe files."""

from __future__ import annotations

import io
import unittest
import zipfile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from test_support import install_aws_stubs

install_aws_stubs()

import resume_docx  # noqa: E402
from resume_store import ResumeError  # noqa: E402


def _docx(*paragraphs: str) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _with_body_xml(source: bytes, fragment: bytes) -> bytes:
    archive = zipfile.ZipFile(io.BytesIO(source))
    xml = archive.read("word/document.xml").replace(b"</w:body>", fragment + b"</w:body>", 1)
    packed = io.BytesIO()
    with zipfile.ZipFile(packed, "w") as out:
        for name in archive.namelist():
            out.writestr(name, xml if name == "word/document.xml" else archive.read(name))
    return packed.getvalue()


def _paragraph_xml(data: bytes, index: int) -> bytes:
    document = Document(io.BytesIO(data))
    paragraphs = [node for node in document.element.body.iter(qn("w:p"))]
    return etree.tostring(paragraphs[index])


class ResumeDocxTests(unittest.TestCase):
    def test_paragraph_ids_follow_body_order(self) -> None:
        mapped = resume_docx.paragraph_map(_docx("Alex Example", "Built tools."))
        self.assertEqual([row["id"] for row in mapped], ["p0", "p1"])
        self.assertEqual(mapped[0]["text"], "Alex Example")

    def test_locked_paragraph_xml_is_unchanged_when_another_is_rewritten(self) -> None:
        source = _docx("Alex Example", "Built tools.")
        untouched = _paragraph_xml(source, 0)
        rewritten, notes = resume_docx.apply(
            source,
            {"roles": []},
            {"replacements": [{"id": "p1", "text": "Built the billing export."}]},
        )
        self.assertEqual(_paragraph_xml(rewritten, 0), untouched)
        self.assertEqual(resume_docx.paragraph_map(rewritten)[1]["text"], "Built the billing export.")
        self.assertEqual(notes, [])

    def test_tracked_changes_are_rejected(self) -> None:
        for fragment in (
            b"<w:delText>old</w:delText>",
            b'<w:ins w:id="1" w:author="Alex Example"/>',
            b'<w:moveFrom w:id="2" w:author="Alex Example"/>',
        ):
            with self.subTest(fragment=fragment):
                with self.assertRaises(ResumeError) as caught:
                    resume_docx.validate(_with_body_xml(_docx("Alex Example"), fragment))
                self.assertIn("Accept all changes", str(caught.exception))

    def test_table_borders_are_not_tracked_changes(self) -> None:
        document = Document()
        document.add_table(rows=2, cols=2).style = "Table Grid"
        buffer = io.BytesIO()
        document.save(buffer)
        borders = b'<w:tblBorders><w:insideH w:val="single"/><w:insideV w:val="single"/></w:tblBorders>'
        resume_docx.validate(_with_body_xml(buffer.getvalue(), borders))
        self.assertIn(
            b"insideH",
            zipfile.ZipFile(io.BytesIO(_with_body_xml(buffer.getvalue(), borders))).read("word/document.xml"),
        )

    def test_macros_are_rejected(self) -> None:
        source = _docx("Alex Example")
        packed = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(source)) as archive, zipfile.ZipFile(packed, "w") as out:
            for name in archive.namelist():
                out.writestr(name, archive.read(name))
            out.writestr("word/vbaProject.bin", b"not-a-macro")
        with self.assertRaises(ResumeError) as caught:
            resume_docx.validate(packed.getvalue())
        self.assertIn("Macro", str(caught.exception))

    def test_custom_entities_are_rejected(self) -> None:
        source = _docx("Alex Example")
        archive = zipfile.ZipFile(io.BytesIO(source))
        xml = archive.read("word/document.xml")
        xml = xml.replace(
            b"<w:document",
            b'<!DOCTYPE x [<!ENTITY x "EXPANDED">]><w:document',
            1,
        )
        packed = io.BytesIO()
        with zipfile.ZipFile(packed, "w") as out:
            for name in archive.namelist():
                out.writestr(name, xml if name == "word/document.xml" else archive.read(name))
        with self.assertRaises(ResumeError) as caught:
            resume_docx.validate(packed.getvalue())
        self.assertIn("not safe", str(caught.exception))
